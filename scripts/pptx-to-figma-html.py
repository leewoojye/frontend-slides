#!/usr/bin/env python3
"""Convert a PPTX into a self-contained, Figma-friendly HTML deck.

The converter intentionally uses a hybrid representation:

* PowerPoint shapes/images/backgrounds are rendered once as an SVG background
  after all editable text is stripped from a temporary copy of the deck.
* Text boxes and PowerPoint placeholders are reconstructed as positioned HTML
  layers and are also embedded in a JSON manifest that the companion Figma
  Slides plugin can import as editable text layers.

This keeps the layout much closer to the source PPTX than the original
``extract-pptx.py`` content-only path while still giving Figma editable text.

Requirements:
    pip install python-pptx
    LibreOffice (``soffice`` or ``libreoffice``)
    Poppler (``pdfseparate`` and ``pdftocairo``)

Usage:
    python scripts/pptx-to-figma-html.py input.pptx --output input.figma.html
"""

from __future__ import annotations

import argparse
import html
import json
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any, Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


CANVAS_WIDTH = 1920.0
CANVAS_HEIGHT = 1080.0
EMU_PER_INCH = 914400.0


@dataclass
class TextLayer:
    name: str
    text: str
    x: float
    y: float
    width: float
    height: float
    fontFamily: str
    fontSize: float
    fontWeight: int
    italic: bool
    color: str
    opacity: float
    align: str
    verticalAlign: str
    rotation: float
    placeholder: bool
    placeholderType: str


def run_checked(command: list[str], cwd: Optional[Path] = None) -> None:
    process = subprocess.run(
        command,
        cwd=str(cwd) if cwd else None,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )
    if process.returncode != 0:
        joined = " ".join(command)
        raise RuntimeError(f"Command failed ({process.returncode}): {joined}\n{process.stdout}")


def find_executable(*names: str) -> str:
    for name in names:
        path = shutil.which(name)
        if path:
            return path
    raise RuntimeError(f"Missing required executable: one of {', '.join(names)}")


def clean_svg(svg: str) -> str:
    svg = re.sub(r"^\s*<\?xml[^>]*>\s*", "", svg, flags=re.I)
    svg = re.sub(r"<!DOCTYPE[^>]*>\s*", "", svg, flags=re.I)
    svg = re.sub(
        r"<svg\b",
        '<svg preserveAspectRatio="none" style="width:100%;height:100%;display:block"',
        svg,
        count=1,
    )
    return svg.strip()


def clear_shape_text(shape: Any) -> None:
    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            clear_shape_text(child)
        return

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                cell.text = ""

    if getattr(shape, "has_text_frame", False):
        shape.text_frame.clear()


def strip_all_text(prs: Presentation) -> None:
    # Master/layout text can be rendered even when the corresponding slide has
    # no explicit shape, so clear all three levels in the render-only copy.
    for master in prs.slide_masters:
        for shape in master.shapes:
            clear_shape_text(shape)
        for layout in master.slide_layouts:
            for shape in layout.shapes:
                clear_shape_text(shape)

    for slide in prs.slides:
        for shape in slide.shapes:
            clear_shape_text(shape)


def render_slides_to_svg(pptx_path: Path, output_dir: Path, slide_count: int) -> list[str]:
    soffice = find_executable("soffice", "libreoffice")
    pdfseparate = find_executable("pdfseparate")
    pdftocairo = find_executable("pdftocairo")

    output_dir.mkdir(parents=True, exist_ok=True)
    run_checked([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(pptx_path)])

    pdf_path = output_dir / f"{pptx_path.stem}.pdf"
    if not pdf_path.exists():
        candidates = sorted(output_dir.glob("*.pdf"), key=lambda p: p.stat().st_mtime, reverse=True)
        if not candidates:
            raise RuntimeError("LibreOffice did not produce a PDF")
        pdf_path = candidates[0]

    result: list[str] = []
    for index in range(1, slide_count + 1):
        single_pdf_pattern = output_dir / "page-%d.pdf"
        run_checked([
            pdfseparate,
            "-f",
            str(index),
            "-l",
            str(index),
            str(pdf_path),
            str(single_pdf_pattern),
        ])
        single_pdf = output_dir / f"page-{index}.pdf"
        svg_path = output_dir / f"slide-{index}.svg"
        run_checked([pdftocairo, "-svg", str(single_pdf), str(svg_path)])
        if not svg_path.exists():
            raise RuntimeError(f"SVG render missing for slide {index}")
        result.append(clean_svg(svg_path.read_text(encoding="utf-8")))

    return result


def enum_name(value: Any) -> str:
    if value is None:
        return ""
    name = getattr(value, "name", None)
    if name:
        return str(name)
    text = str(value)
    return re.split(r"[\s(]", text, maxsplit=1)[0].upper()


def rgb_to_hex(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip().replace("#", "")
    if re.fullmatch(r"[0-9A-Fa-f]{6}", text):
        return f"#{text.upper()}"
    return None


def safe_run_color(run: Any) -> Optional[str]:
    try:
        return rgb_to_hex(run.font.color.rgb)
    except Exception:
        return None


def first_run(shape: Any) -> Any:
    if not getattr(shape, "has_text_frame", False):
        return None
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text or run.font.name or run.font.size:
                return run
    return None


def first_paragraph(shape: Any) -> Any:
    if not getattr(shape, "has_text_frame", False):
        return None
    paragraphs = shape.text_frame.paragraphs
    return paragraphs[0] if paragraphs else None


def placeholder_prompt(type_name: str) -> str:
    key = type_name.upper()
    if "TITLE" in key:
        return "Click to add title"
    if "SUBTITLE" in key:
        return "Click to add subtitle"
    if any(token in key for token in ("PICTURE", "MEDIA", "CLIP_ART")):
        return "Click to add image"
    if any(token in key for token in ("CHART", "TABLE", "OBJECT", "CONTENT")):
        return "Click to add content"
    if any(token in key for token in ("DATE", "FOOTER", "SLIDE_NUMBER", "HEADER")):
        return ""
    return "Click to add text"


def placeholder_default_font_size(type_name: str) -> float:
    key = type_name.upper()
    if "TITLE" in key:
        return 64.0
    if "SUBTITLE" in key:
        return 34.0
    return 28.0


def alignment_name(value: Any) -> str:
    name = enum_name(value)
    if name in {"CENTER", "DISTRIBUTE", "THAI_DISTRIBUTE"}:
        return "center"
    if name == "RIGHT":
        return "right"
    if name == "JUSTIFY":
        return "justify"
    return "left"


def vertical_alignment_name(value: Any) -> str:
    name = enum_name(value)
    if name in {"MIDDLE", "CENTER"}:
        return "center"
    if name == "BOTTOM":
        return "bottom"
    return "top"


def shape_placeholder_index(shape: Any) -> Optional[int]:
    try:
        if shape.is_placeholder:
            return int(shape.placeholder_format.idx)
    except Exception:
        pass
    return None


def shape_placeholder_type(shape: Any) -> str:
    try:
        if shape.is_placeholder:
            return enum_name(shape.placeholder_format.type)
    except Exception:
        pass
    return ""


def find_layout_placeholder(slide: Any, idx: Optional[int]) -> Any:
    if idx is None:
        return None
    try:
        for ph in slide.slide_layout.placeholders:
            if int(ph.placeholder_format.idx) == idx:
                return ph
    except Exception:
        pass
    return None


def maybe_attr(primary: Any, fallback: Any, attr: str, default: Any = None) -> Any:
    value = getattr(primary, attr, None) if primary is not None else None
    if value is not None:
        return value
    value = getattr(fallback, attr, None) if fallback is not None else None
    return default if value is None else value


def shape_geometry(shape: Any, fallback: Any, sx: float, sy: float) -> tuple[float, float, float, float]:
    left = maybe_attr(shape, fallback, "left", 0)
    top = maybe_attr(shape, fallback, "top", 0)
    width = maybe_attr(shape, fallback, "width", 0)
    height = maybe_attr(shape, fallback, "height", 0)
    return float(left) * sx, float(top) * sy, float(width) * sx, float(height) * sy


def font_px_from_run(run: Any, fallback_run: Any, px_per_inch: float, default_px: float) -> float:
    size = maybe_attr(getattr(run, "font", None), getattr(fallback_run, "font", None), "size")
    if size is None:
        return default_px
    try:
        return max(1.0, float(size.pt) * px_per_inch / 72.0)
    except Exception:
        return default_px


def font_name_from_run(run: Any, fallback_run: Any, default_font: str) -> str:
    name = maybe_attr(getattr(run, "font", None), getattr(fallback_run, "font", None), "name")
    return str(name).strip() if name else default_font


def font_weight_from_run(run: Any, fallback_run: Any, placeholder: bool) -> int:
    bold = maybe_attr(getattr(run, "font", None), getattr(fallback_run, "font", None), "bold")
    if bold is True:
        return 700
    if bold is False:
        return 400
    return 600 if placeholder else 400


def italic_from_run(run: Any, fallback_run: Any) -> bool:
    italic = maybe_attr(getattr(run, "font", None), getattr(fallback_run, "font", None), "italic")
    return bool(italic) if italic is not None else False


def color_from_run(run: Any, fallback_run: Any, placeholder: bool) -> tuple[str, float]:
    color = safe_run_color(run) if run is not None else None
    if not color and fallback_run is not None:
        color = safe_run_color(fallback_run)
    if placeholder:
        return color or "#777777", 0.58
    return color or "#202124", 1.0


def text_for_shape(shape: Any, placeholder_type: str) -> tuple[str, bool]:
    raw = ""
    if getattr(shape, "has_text_frame", False):
        raw = shape.text or ""
    if raw.strip():
        return raw, False
    if getattr(shape, "is_placeholder", False):
        prompt = placeholder_prompt(placeholder_type)
        if prompt:
            return prompt, True
    return "", False


def collect_text_shape(
    shape: Any,
    slide: Any,
    sx: float,
    sy: float,
    px_per_inch: float,
    source_prefix: str,
    default_font: str,
) -> Optional[TextLayer]:
    if not getattr(shape, "has_text_frame", False):
        return None

    ph_idx = shape_placeholder_index(shape)
    ph_type = shape_placeholder_type(shape)
    fallback = find_layout_placeholder(slide, ph_idx) if source_prefix == "slide" else None

    text, is_prompt = text_for_shape(shape, ph_type)
    if not text:
        return None

    x, y, width, height = shape_geometry(shape, fallback, sx, sy)
    if width <= 0 or height <= 0:
        return None

    run = first_run(shape)
    fallback_run = first_run(fallback) if fallback is not None else None
    default_px = placeholder_default_font_size(ph_type) if is_prompt else 28.0
    font_size = font_px_from_run(run, fallback_run, px_per_inch, default_px)
    font_family = font_name_from_run(run, fallback_run, default_font)
    font_weight = font_weight_from_run(run, fallback_run, is_prompt)
    italic = italic_from_run(run, fallback_run)
    color, opacity = color_from_run(run, fallback_run, is_prompt)

    text_frame = shape.text_frame
    fallback_tf = fallback.text_frame if fallback is not None and getattr(fallback, "has_text_frame", False) else None

    # Apply PowerPoint internal text-box margins directly to the editable area.
    ml = maybe_attr(text_frame, fallback_tf, "margin_left", 0) or 0
    mr = maybe_attr(text_frame, fallback_tf, "margin_right", 0) or 0
    mt = maybe_attr(text_frame, fallback_tf, "margin_top", 0) or 0
    mb = maybe_attr(text_frame, fallback_tf, "margin_bottom", 0) or 0
    x += float(ml) * sx
    y += float(mt) * sy
    width = max(1.0, width - (float(ml) + float(mr)) * sx)
    height = max(1.0, height - (float(mt) + float(mb)) * sy)

    paragraph = first_paragraph(shape)
    fallback_paragraph = first_paragraph(fallback) if fallback is not None else None
    align_value = maybe_attr(paragraph, fallback_paragraph, "alignment")
    vertical_value = maybe_attr(text_frame, fallback_tf, "vertical_anchor")

    rotation = float(getattr(shape, "rotation", 0.0) or 0.0)
    name = f"{source_prefix}:{getattr(shape, 'name', 'text')}"

    return TextLayer(
        name=name,
        text=text,
        x=round(x, 3),
        y=round(y, 3),
        width=round(width, 3),
        height=round(height, 3),
        fontFamily=font_family,
        fontSize=round(font_size, 3),
        fontWeight=font_weight,
        italic=italic,
        color=color,
        opacity=opacity,
        align=alignment_name(align_value),
        verticalAlign=vertical_alignment_name(vertical_value),
        rotation=rotation,
        placeholder=is_prompt,
        placeholderType=ph_type,
    )


def collect_group_text(
    shape: Any,
    slide: Any,
    sx: float,
    sy: float,
    px_per_inch: float,
    source_prefix: str,
    default_font: str,
) -> list[TextLayer]:
    # Group text is uncommon in presentation templates. python-pptx exposes
    # child geometry relative to the group, so complicated scaled groups may
    # still need manual adjustment. The SVG remains authoritative for visuals.
    layers: list[TextLayer] = []
    for child in shape.shapes:
        if getattr(child, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            layers.extend(collect_group_text(child, slide, sx, sy, px_per_inch, source_prefix, default_font))
        else:
            layer = collect_text_shape(child, slide, sx, sy, px_per_inch, source_prefix, default_font)
            if layer:
                layers.append(layer)
    return layers


def collect_slide_text_layers(prs: Presentation, slide: Any, default_font: str) -> list[TextLayer]:
    sx = CANVAS_WIDTH / float(prs.slide_width)
    sy = CANVAS_HEIGHT / float(prs.slide_height)
    slide_width_inches = float(prs.slide_width) / EMU_PER_INCH
    px_per_inch = CANVAS_WIDTH / slide_width_inches

    layers: list[TextLayer] = []
    slide_ph_idxs = {
        idx for idx in (shape_placeholder_index(s) for s in slide.shapes) if idx is not None
    }

    master = slide.slide_layout.slide_master
    for shape in master.shapes:
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            layers.extend(collect_group_text(shape, slide, sx, sy, px_per_inch, "master", default_font))
            continue
        if getattr(shape, "is_placeholder", False) and not (getattr(shape, "text", "") or "").strip():
            continue
        layer = collect_text_shape(shape, slide, sx, sy, px_per_inch, "master", default_font)
        if layer:
            layers.append(layer)

    for shape in slide.slide_layout.shapes:
        idx = shape_placeholder_index(shape)
        if idx is not None and idx in slide_ph_idxs:
            continue
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            layers.extend(collect_group_text(shape, slide, sx, sy, px_per_inch, "layout", default_font))
            continue
        layer = collect_text_shape(shape, slide, sx, sy, px_per_inch, "layout", default_font)
        if layer:
            layers.append(layer)

    for shape in slide.shapes:
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            layers.extend(collect_group_text(shape, slide, sx, sy, px_per_inch, "slide", default_font))
            continue
        layer = collect_text_shape(shape, slide, sx, sy, px_per_inch, "slide", default_font)
        if layer:
            layers.append(layer)

    return layers


def layer_style(layer: TextLayer) -> str:
    weight = layer.fontWeight
    italic = "italic" if layer.italic else "normal"
    justify = {"top": "flex-start", "center": "center", "bottom": "flex-end"}.get(layer.verticalAlign, "flex-start")
    safe_family = layer.fontFamily.replace('"', "'")
    return (
        f"left:{layer.x}px;top:{layer.y}px;width:{layer.width}px;height:{layer.height}px;"
        f"font-family:\"{safe_family}\",Arial,sans-serif;font-size:{layer.fontSize}px;"
        f"font-weight:{weight};font-style:{italic};color:{layer.color};opacity:{layer.opacity};"
        f"text-align:{layer.align};justify-content:{justify};"
        f"transform:rotate({layer.rotation}deg);"
    )


def escape_script_json(value: Any) -> str:
    text = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return text.replace("</script>", "<\\/script>")


def build_html(deck_name: str, svgs: list[str], text_layers: list[list[TextLayer]]) -> str:
    slides_manifest = []
    slide_sections = []

    for index, (svg, layers) in enumerate(zip(svgs, text_layers), start=1):
        layer_dicts = [asdict(layer) for layer in layers]
        slides_manifest.append({
            "name": f"Slide {index}",
            "svg": svg,
            "textLayers": layer_dicts,
        })

        overlays = []
        for layer_index, layer in enumerate(layers):
            classes = "ppt-text placeholder" if layer.placeholder else "ppt-text"
            overlays.append(
                f'<div class="{classes}" data-layer-index="{layer_index}" '
                f'data-layer-name="{html.escape(layer.name, quote=True)}" '
                f'style="{html.escape(layer_style(layer), quote=True)}">'
                f'<span>{html.escape(layer.text).replace(chr(10), "<br>")}</span></div>'
            )
        slide_sections.append(
            f'<section class="slide{" active" if index == 1 else ""}" data-slide-index="{index - 1}">'
            f'<div class="ppt-vector-bg">{svg}</div>'
            f'{"".join(overlays)}'
            f'</section>'
        )

    manifest = {
        "version": 1,
        "generator": "frontend-slides pptx-to-figma-html",
        "name": deck_name,
        "width": CANVAS_WIDTH,
        "height": CANVAS_HEIGHT,
        "slides": slides_manifest,
    }

    return f'''<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{html.escape(deck_name)}</title>
<style>
  * {{ box-sizing: border-box; }}
  html, body {{ margin: 0; width: 100%; min-height: 100%; background: #171719; }}
  body {{ overflow: hidden; color: #fff; font-family: Arial, sans-serif; }}
  .viewport {{ position: fixed; inset: 0; overflow: hidden; }}
  .stage {{ position: absolute; width: {CANVAS_WIDTH}px; height: {CANVAS_HEIGHT}px; transform-origin: 0 0; }}
  .slide {{ position: absolute; inset: 0; width: {CANVAS_WIDTH}px; height: {CANVAS_HEIGHT}px; overflow: hidden; background: #fff; visibility: hidden; opacity: 0; }}
  .slide.active {{ visibility: visible; opacity: 1; }}
  .ppt-vector-bg {{ position: absolute; inset: 0; width: 100%; height: 100%; }}
  .ppt-vector-bg > svg {{ width: 100% !important; height: 100% !important; display: block; }}
  .ppt-text {{ position: absolute; z-index: 2; display: flex; overflow: hidden; white-space: pre-wrap; line-height: 1.12; transform-origin: center center; pointer-events: auto; }}
  .ppt-text > span {{ width: 100%; align-self: inherit; }}
  .ppt-text.placeholder {{ outline: 2px dashed rgba(90,90,90,.35); outline-offset: -2px; background: rgba(255,255,255,.035); }}
  .controls {{ position: fixed; z-index: 20; left: 50%; bottom: 18px; transform: translateX(-50%); display: flex; align-items: center; gap: 10px; padding: 7px 10px; border-radius: 999px; background: rgba(15,15,18,.78); backdrop-filter: blur(10px); }}
  .controls button {{ border: 0; border-radius: 8px; padding: 8px 12px; color: #fff; background: #34343a; cursor: pointer; }}
  .controls button:hover {{ background: #484850; }}
  #counter {{ min-width: 72px; text-align: center; font-size: 13px; }}
  body.figma-export {{ overflow: auto; background: #29292d; }}
  body.figma-export .viewport {{ position: static; overflow: visible; }}
  body.figma-export .stage {{ position: static; width: {CANVAS_WIDTH}px; height: auto; transform: none !important; margin: 40px auto; }}
  body.figma-export .slide {{ position: relative; inset: auto; display: block; visibility: visible; opacity: 1; margin: 0 0 64px 0; box-shadow: 0 18px 50px rgba(0,0,0,.28); }}
  body.figma-export .controls {{ display: none; }}
  @media print {{
    body {{ overflow: visible; background: #fff; }}
    .viewport, .stage {{ position: static; transform: none !important; width: auto; height: auto; }}
    .slide {{ position: relative; visibility: visible; opacity: 1; break-after: page; page-break-after: always; }}
    .controls {{ display: none; }}
  }}
</style>
</head>
<body>
<div class="viewport"><main class="stage" id="stage">{''.join(slide_sections)}</main></div>
<nav class="controls">
  <button type="button" id="prev">←</button>
  <span id="counter">1 / {len(svgs)}</span>
  <button type="button" id="next">→</button>
  <button type="button" id="exportMode">Figma view</button>
  <button type="button" id="downloadManifest">Manifest JSON</button>
</nav>
<script id="figma-deck-manifest" type="application/json">{escape_script_json(manifest)}</script>
<script>
(() => {{
  const slides = [...document.querySelectorAll('.slide')];
  const stage = document.getElementById('stage');
  const counter = document.getElementById('counter');
  let current = 0;

  function fit() {{
    if (document.body.classList.contains('figma-export')) return;
    const scale = Math.min(innerWidth / {CANVAS_WIDTH}, innerHeight / {CANVAS_HEIGHT});
    const x = (innerWidth - {CANVAS_WIDTH} * scale) / 2;
    const y = (innerHeight - {CANVAS_HEIGHT} * scale) / 2;
    stage.style.transform = `translate(${{x}}px, ${{y}}px) scale(${{scale}})`;
  }}

  function show(index) {{
    current = Math.max(0, Math.min(index, slides.length - 1));
    slides.forEach((slide, i) => slide.classList.toggle('active', i === current));
    counter.textContent = `${{current + 1}} / ${{slides.length}}`;
    history.replaceState(null, '', `#slide-${{current + 1}}`);
  }}

  document.getElementById('prev').onclick = () => show(current - 1);
  document.getElementById('next').onclick = () => show(current + 1);
  document.addEventListener('keydown', (event) => {{
    if (['ArrowRight', 'ArrowDown', 'PageDown', ' '].includes(event.key)) {{ event.preventDefault(); show(current + 1); }}
    if (['ArrowLeft', 'ArrowUp', 'PageUp'].includes(event.key)) {{ event.preventDefault(); show(current - 1); }}
  }});

  document.getElementById('exportMode').onclick = () => {{
    document.body.classList.toggle('figma-export');
    stage.style.transform = 'none';
    fit();
  }};

  document.getElementById('downloadManifest').onclick = () => {{
    const data = document.getElementById('figma-deck-manifest').textContent;
    const blob = new Blob([data], {{ type: 'application/json' }});
    const url = URL.createObjectURL(blob);
    const link = document.createElement('a');
    link.href = url;
    link.download = '{html.escape(deck_name)}.figma.json';
    link.click();
    URL.revokeObjectURL(url);
  }};

  const initial = location.hash.match(/slide-(\\d+)/);
  if (new URLSearchParams(location.search).get('figma') === '1') document.body.classList.add('figma-export');
  show(initial ? Number(initial[1]) - 1 : 0);
  fit();
  addEventListener('resize', fit);
}})();
</script>
</body>
</html>
'''


def convert(input_path: Path, output_path: Path, keep_workdir: Optional[Path], default_font: str) -> None:
    prs = Presentation(str(input_path))
    slide_count = len(prs.slides)
    if slide_count == 0:
        raise RuntimeError("The presentation has no slides")

    all_text_layers = [collect_slide_text_layers(prs, slide, default_font) for slide in prs.slides]

    if keep_workdir:
        workdir = keep_workdir
        workdir.mkdir(parents=True, exist_ok=True)
        cleanup = False
    else:
        temp = tempfile.TemporaryDirectory(prefix="frontend-slides-figma-")
        workdir = Path(temp.name)
        cleanup = True

    try:
        render_copy = Presentation(str(input_path))
        strip_all_text(render_copy)
        stripped_pptx = workdir / f"{input_path.stem}.text-stripped.pptx"
        render_copy.save(str(stripped_pptx))

        svgs = render_slides_to_svg(stripped_pptx, workdir / "render", slide_count)
        document = build_html(input_path.stem, svgs, all_text_layers)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(document, encoding="utf-8")

        print(f"Wrote {output_path}")
        print(f"Slides: {slide_count}")
        print(f"Editable text layers: {sum(len(x) for x in all_text_layers)}")
        print("The output is self-contained; images are embedded in the inline SVG backgrounds.")
    finally:
        if cleanup:
            temp.cleanup()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Input .pptx file")
    parser.add_argument("--output", "-o", type=Path, help="Output standalone HTML file")
    parser.add_argument("--work-dir", type=Path, help="Keep intermediate stripped PPTX/PDF/SVG files here")
    parser.add_argument(
        "--default-font",
        default="Noto Sans KR",
        help="Fallback font when the PPTX only references a theme font (default: Noto Sans KR)",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path = args.input.resolve()
    if not input_path.exists():
        print(f"Input not found: {input_path}", file=sys.stderr)
        return 2
    if input_path.suffix.lower() != ".pptx":
        print("Input must be a .pptx file", file=sys.stderr)
        return 2

    output_path = (args.output or input_path.with_suffix(".figma.html")).resolve()
    try:
        convert(input_path, output_path, args.work_dir.resolve() if args.work_dir else None, args.default_font)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
