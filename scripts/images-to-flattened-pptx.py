#!/usr/bin/env python3
from pathlib import Path
import argparse, re
from pptx import Presentation
from pptx.util import Inches


def natural_key(path: Path):
    return [int(x) if x.isdigit() else x.lower() for x in re.split(r'(\d+)', path.name)]


def main():
    p = argparse.ArgumentParser()
    p.add_argument('image_dir', type=Path)
    p.add_argument('output', type=Path)
    args = p.parse_args()

    images = sorted([x for x in args.image_dir.iterdir() if x.suffix.lower() in {'.png','.jpg','.jpeg'}], key=natural_key)
    if not images:
        raise SystemExit('No slide images found')

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # remove the initial default slide if present in this python-pptx version
    while len(prs.slides):
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]

    for img in images:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.output)
    print(f'Created {args.output} with {len(images)} image-only slides')

if __name__ == '__main__':
    main()
